"""Download and parse explicitly dated public monthly input-cost series."""

from __future__ import annotations

import argparse
import csv
import hashlib
import json
import math
import re
import urllib.request
from collections import defaultdict
from datetime import date, datetime, timedelta
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]
METALS = {"copper": "Copper", "aluminium": "Aluminum", "tin": "Tin"}


def next_month(day):
    return date(day.year + (day.month == 12), day.month % 12 + 1, 1)


def validate_months(rows):
    dates = [date.fromisoformat(r["date"]) for r in rows]
    if not dates or any(d.day != 1 for d in dates):
        raise ValueError("Monthly observations require first-of-month period labels")
    if any(right != next_month(left) for left, right in zip(dates, dates[1:])):
        raise ValueError("Missing, duplicate, or unsorted months; no imputation is permitted")
    if any(not math.isfinite(r["value"]) or r["value"] <= 0 for r in rows):
        raise ValueError("Observed cost must be finite and positive")


def parse_metals(path, cutoff):
    from openpyxl import load_workbook
    workbook = load_workbook(path, read_only=True, data_only=True)
    try:
        rows = workbook["Monthly Prices"].iter_rows(values_only=True)
        header = None
        for row in rows:
            if set(METALS.values()) <= set(row):
                header = list(row)
                break
        if header is None:
            raise ValueError("Pink Sheet metal columns not found")
        units = next(rows)
        columns = {key: header.index(label) for key, label in METALS.items()}
        if any(units[c] != "($/mt)" for c in columns.values()):
            raise ValueError("Unexpected metal units; refusing an implicit conversion")
        result = {key: [] for key in METALS}
        for row in rows:
            if not isinstance(row[0], str) or not re.fullmatch(r"\d{4}M\d{2}", row[0]):
                continue
            month = row[0][:4] + "-" + row[0][5:] + "-01"
            if month > cutoff:
                continue
            for key, column in columns.items():
                value = row[column]
                if type(value) not in {int, float}:
                    raise ValueError(f"{key} {month}: missing/invalid price")
                result[key].append({"date": month, "value": float(value)})
        for history in result.values():
            validate_months(history)
        return result
    finally:
        workbook.close()


def parse_freight(path, cutoff):
    groups, observed = defaultdict(list), []
    with path.open(encoding="utf-8-sig", newline="") as handle:
        for row in csv.DictReader(handle):
            day = datetime.strptime(row["Date"], "%Y/%m/%d").date()
            observed.append(day)
            month = day.replace(day=1).isoformat()
            if month <= cutoff:
                groups[month].append(float(row["SCFI_EU"]))
    if observed != sorted(set(observed)):
        raise ValueError("Freight dates must be unique and sorted")
    # Discard an incomplete boundary month instead of inventing absent quotes.
    if observed[0].day > 7:
        groups.pop(observed[0].replace(day=1).isoformat(), None)
    if observed[-1] + timedelta(days=7) < next_month(observed[-1]):
        groups.pop(observed[-1].replace(day=1).isoformat(), None)
    history = [{"date": month, "value": sum(values) / len(values), "weeklyObservations": len(values)} for month, values in sorted(groups.items())]
    validate_months(history)
    return history


def prepare(download=False, root=ROOT):
    registry = json.loads((root / "data/sources/costs.json").read_text(encoding="utf-8"))
    raw = root / "data/raw/costs"
    raw.mkdir(parents=True, exist_ok=True)
    for source in registry["downloads"]:
        path = raw / source["file"]
        if not path.exists() and download:
            request = urllib.request.Request(source["url"], headers={"User-Agent": "Ripple public-data research"})
            with urllib.request.urlopen(request, timeout=60) as response:
                path.write_bytes(response.read())
        if hashlib.sha256(path.read_bytes()).hexdigest() != source["sha256"]:
            raise ValueError(f"Changed public download: {source['file']}; review and update the registry explicitly")
    cutoff = registry["lastCompleteMonth"]
    metals = parse_metals(raw / registry["downloads"][0]["file"], cutoff)
    series = []
    for key, history in metals.items():
        series.append({"id": key, "label": key.capitalize(), "unit": "USD per tonne", "history": history,
                       "source": {k: registry["downloads"][0][k] for k in ["title", "url"]}})
    freight = registry["downloads"][1]
    history = parse_freight(raw / freight["file"], cutoff)
    series.append({"id": "container-freight", "label": "Container freight: Shanghai–Europe (archive through April 2025)",
                   "unit": "USD per TEU", "history": history,
                   "source": {k: freight[k] for k in ["title", "url"]}, "definitionSource": registry["freightDefinition"],
                   "availabilityNote": "Archived public research data end April 2025. Forecasts start May 2025; these are not current freight quotes."})
    for item in series:
        item["historyThrough"] = item["history"][-1]["date"]
        item["forecastStatus"] = "archived" if item["historyThrough"] < cutoff else "latest verified vintage"
    document = {"sampleData": False, "asOf": registry["asOf"], "series": series}
    path = root / "data/processed/cost_history.json"
    path.write_text(json.dumps(document, ensure_ascii=False, separators=(",", ":"), allow_nan=False) + "\n", encoding="utf-8")
    (root / "data/review/cost_gaps.json").write_text(json.dumps({
        "sampleData": False, "seriesId": "container-freight", "lastObservedMonth": series[-1]["historyThrough"],
        "latestCompleteMonth": cutoff, "reason": "Verified open research supplement ends April 2025; no current history was substituted",
        "handling": "Explicit archived forecast from May 2025, not current freight pricing or a filled history",
        "source": series[-1]["source"],
    }, indent=2) + "\n", encoding="utf-8")
    for item in series:
        print(item["id"], len(item["history"]), item["history"][0]["date"], item["historyThrough"])
    return document


if __name__ == "__main__":
    parser = argparse.ArgumentParser(description=__doc__)
    parser.add_argument("--download", action="store_true")
    prepare(parser.parse_args().download)
